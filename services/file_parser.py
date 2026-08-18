"""文件解析服务 - Excel/CSV 读取"""
from pathlib import Path
import csv
import io
from openpyxl import load_workbook


class FileParserService:
    """文件解析服务"""

    def parse_file(self, file_path: str) -> str:
        """根据文件扩展名自动解析"""
        ext = Path(file_path).suffix.lower()
        if ext in (".xlsx", ".xls"):
            return self.parse_excel(file_path)
        elif ext in (".csv", ".tsv"):
            return self.parse_csv(file_path)
        else:
            raise ValueError(f"不支持的文件格式: {ext}")

    def parse_excel(self, file_path: str) -> str:
        """解析 Excel 文件为 CSV 字符串"""
        wb = load_workbook(file_path, read_only=True, data_only=True)
        ws = wb.active
        if ws is None:
            return ""

        output = io.StringIO()
        writer = csv.writer(output)
        for row in ws.iter_rows(values_only=True):
            writer.writerow([str(cell) if cell is not None else "" for cell in row])

        wb.close()
        return output.getvalue()

    def parse_csv(self, file_path: str) -> str:
        """读取 CSV 文件"""
        with open(file_path, "r", encoding="utf-8-sig") as f:
            return f.read()
